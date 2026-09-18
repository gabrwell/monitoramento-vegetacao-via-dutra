"""Gera o relatório acadêmico final em DOCX a partir dos resultados do projeto."""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "relatorio" / "relatorio_final.docx"
AZUL = "1F4E78"
AZUL_CLARO = "EAF2F8"
CINZA = "D9D9D9"
BRANCO = "FFFFFF"
PRETO = RGBColor(0, 0, 0)
INTEGRANTES = [
    ("Pedro Henrique dos Santos Cardoso", "563268", "2CCPG"),
    ("Gabriel Gibin Leoncio", "565462", "2CCPG"),
    ("Rafael do Nascimento Silva", "566263", "2CCPG"),
    ("Rai Augusto Ribeiro", "562870", "2CCPG"),
    ("Guilherme Morais de Assis", "564198", "2CCPG"),
    ("Lucas Werpp Franco", "556044", "2CCPG"),
]


def fonte(run, nome="Arial", tamanho=11, negrito=False, italico=False, cor=PRETO):
    run.font.name = nome
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), nome)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), nome)
    run.font.size = Pt(tamanho)
    run.font.bold = negrito
    run.font.italic = italico
    run.font.color.rgb = cor
    return run


def sombreamento(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def bordas_tabela(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for nome in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{nome}")
        edge = borders.find(tag)
        if edge is None:
            edge = OxmlElement(f"w:{nome}")
            borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "6")
        edge.set(qn("w:color"), CINZA)


def margem_celula(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for nome, valor in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        elemento = tc_mar.find(qn(f"w:{nome}"))
        if elemento is None:
            elemento = OxmlElement(f"w:{nome}")
            tc_mar.append(elemento)
        elemento.set(qn("w:w"), str(valor))
        elemento.set(qn("w:type"), "dxa")


def repetir_cabecalho(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def tabela(doc, cabecalho, linhas, larguras=None, alinhamentos=None):
    table = doc.add_table(rows=1, cols=len(cabecalho))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    repetir_cabecalho(table.rows[0])
    for indice, texto in enumerate(cabecalho):
        cell = table.rows[0].cells[indice]
        sombreamento(cell, AZUL)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        margem_celula(cell)
        paragrafo = cell.paragraphs[0]
        paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragrafo.paragraph_format.space_after = Pt(0)
        fonte(paragrafo.add_run(str(texto)), tamanho=9, negrito=True, cor=RGBColor(255, 255, 255))
    for numero, linha in enumerate(linhas):
        cells = table.add_row().cells
        for indice, valor in enumerate(linha):
            cell = cells[indice]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            margem_celula(cell)
            if numero % 2:
                sombreamento(cell, AZUL_CLARO)
            paragrafo = cell.paragraphs[0]
            paragrafo.alignment = (
                alinhamentos[indice] if alinhamentos else WD_ALIGN_PARAGRAPH.LEFT
            )
            paragrafo.paragraph_format.space_after = Pt(0)
            fonte(paragrafo.add_run(str(valor)), tamanho=9)
    if larguras:
        for row in table.rows:
            for indice, largura in enumerate(larguras):
                row.cells[indice].width = Inches(largura)
    bordas_tabela(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def adicionar_paragrafo(doc, texto, estilo=None, negrito_inicial=None):
    p = doc.add_paragraph(style=estilo)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.12
    if negrito_inicial and texto.startswith(negrito_inicial):
        fonte(p.add_run(negrito_inicial), negrito=True)
        fonte(p.add_run(texto[len(negrito_inicial) :]))
    else:
        fonte(p.add_run(texto))
    return p


def adicionar_lista(doc, itens):
    for item in itens:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        fonte(p.add_run(item))


def adicionar_figura(doc, caminho, legenda, largura=6.35):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    imagem = p.add_run().add_picture(str(caminho), width=Inches(largura))
    imagem._inline.docPr.set("title", legenda)
    imagem._inline.docPr.set("descr", legenda)
    legenda_p = doc.add_paragraph()
    legenda_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    legenda_p.paragraph_format.space_after = Pt(8)
    fonte(legenda_p.add_run(legenda), tamanho=9, italico=True)


def adicionar_equacao_ndvi(doc):
    xml = f"""
    <m:oMathPara {nsdecls('m')}>
      <m:oMath>
        <m:r><m:t>NDVI = </m:t></m:r>
        <m:f>
          <m:num><m:r><m:t>B08 - B04</m:t></m:r></m:num>
          <m:den><m:r><m:t>B08 + B04</m:t></m:r></m:den>
        </m:f>
      </m:oMath>
    </m:oMathPara>
    """
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p._p.append(parse_xml(xml))
    p.paragraph_format.space_after = Pt(8)


def adicionar_numero_pagina(paragrafo):
    paragrafo.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fonte(paragrafo.add_run("Página "), tamanho=9, cor=RGBColor(90, 90, 90))
    partes = (("begin", None), (None, " PAGE "), ("separate", None))
    for tipo, instrucao in partes:
        run = paragrafo.add_run()
        if tipo:
            campo = OxmlElement("w:fldChar")
            campo.set(qn("w:fldCharType"), tipo)
            run._r.append(campo)
        else:
            texto = OxmlElement("w:instrText")
            texto.set(qn("xml:space"), "preserve")
            texto.text = instrucao
            run._r.append(texto)
        fonte(run, tamanho=9, cor=RGBColor(90, 90, 90))
    fonte(paragrafo.add_run("1"), tamanho=9, cor=RGBColor(90, 90, 90))
    run = paragrafo.add_run()
    campo = OxmlElement("w:fldChar")
    campo.set(qn("w:fldCharType"), "end")
    run._r.append(campo)
    fonte(run, tamanho=9, cor=RGBColor(90, 90, 90))


def main():
    resumo = json.loads((ROOT / "outputs/resumo_resultados.json").read_text(encoding="utf-8"))
    segmentos = pd.read_csv(ROOT / "data/processed/segmentos_priorizados.csv")

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.orientation = WD_ORIENT.PORTRAIT
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(11)
    normal.font.color.rgb = PRETO

    titulo_style = doc.styles["Title"]
    titulo_style.font.name = "Arial"
    titulo_style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    titulo_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    titulo_style.font.size = Pt(23)
    titulo_style.font.bold = True
    titulo_style.font.color.rgb = PRETO

    for nome, tamanho in (("Heading 1", 15), ("Heading 2", 12.5)):
        estilo = doc.styles[nome]
        estilo.font.name = "Arial"
        estilo._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        estilo._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        estilo.font.size = Pt(tamanho)
        estilo.font.bold = True
        estilo.font.color.rgb = PRETO
        estilo.paragraph_format.space_before = Pt(12)
        estilo.paragraph_format.space_after = Pt(6)
        estilo.paragraph_format.keep_with_next = True

    adicionar_numero_pagina(section.footer.paragraphs[0])
    doc.core_properties.title = "Priorização de inspeções de vegetação na Via Dutra"
    doc.core_properties.subject = "Prova de conceito de Ciência de Dados com Sentinel 2"
    doc.core_properties.keywords = "Sentinel-2, NDVI, rodovias, vegetação, Ciência de Dados"
    doc.core_properties.author = "Turma 2CCPG"

    # Capa
    doc.add_paragraph().paragraph_format.space_after = Pt(55)
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    fonte(p.add_run("Priorização de inspeções de vegetação na Via Dutra"), tamanho=23, negrito=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(22)
    fonte(p.add_run("Prova de conceito de Ciência de Dados com Sentinel 2"), tamanho=14)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fonte(p.add_run("BR 116 entre Jacareí e São José dos Campos"), tamanho=12, negrito=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fonte(p.add_run("CCR RioSP Motiva"), tamanho=11)
    p.paragraph_format.space_after = Pt(12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(5)
    fonte(p.add_run("Integrantes da turma 2CCPG"), tamanho=11, negrito=True)
    for nome, rm, turma in INTEGRANTES:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1)
        fonte(p.add_run(f"{nome} - RM {rm} - {turma}"), tamanho=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(38)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fonte(p.add_run("Setembro de 2026"), tamanho=11)
    doc.add_page_break()

    # Resumo e principais resultados
    doc.add_heading("Resumo", level=1)
    adicionar_paragrafo(
        doc,
        "Este trabalho apresenta uma prova de conceito para priorizar inspeções de vegetação "
        "na BR-116 entre Jacareí e São José dos Campos. O pipeline combina o traçado público "
        "da rodovia com cenas Sentinel-2 L2A de agosto de 2024 e agosto de 2025. Foram "
        "construídas 184 observações reais em 92 pontos, classificadas por uma regra proxy "
        "explicável. O resultado indica 65 pontos de prioridade baixa, 22 média e 5 alta. "
        "As classes organizam a inspeção e não constituem laudo de risco.",
    )
    doc.add_heading("Principais resultados", level=2)
    tabela(
        doc,
        ["Indicador", "Resultado"],
        [
            ["Observações tabulares reais", "184"],
            ["Pontos comparáveis", "92"],
            ["Períodos", "Agosto de 2024 e agosto de 2025"],
            ["Prioridade baixa", "65 pontos 70,7 por cento"],
            ["Prioridade média", "22 pontos 23,9 por cento"],
            ["Prioridade alta", "5 pontos 5,4 por cento"],
            ["Variação média do NDVI", "-0,0064"],
        ],
        larguras=[3.4, 3.2],
        alinhamentos=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER],
    )
    adicionar_paragrafo(
        doc,
        "Conclusão principal: o processo automatizado consegue gerar uma fila inicial de "
        "inspeção com dados públicos reais. A resolução espacial e a ausência de validação "
        "em campo impedem interpretar essa fila como diagnóstico de segurança.",
        negrito_inicial="Conclusão principal:",
    )

    doc.add_heading("1 Introdução", level=1)
    adicionar_paragrafo(
        doc,
        "A gestão da vegetação ao longo de rodovias exige acompanhamento contínuo para "
        "planejar inspeções e intervenções. Como o desafio não prevê uma campanha própria "
        "de coleta em campo, foram usadas fontes públicas e APIs. Além do resultado, o "
        "trabalho registra as decisões metodológicas, a qualidade dos dados e as limitações.",
    )
    adicionar_paragrafo(
        doc,
        "O projeto segue o caminho tabular previsto no enunciado: fonte pública, coleta, "
        "tratamento, análise exploratória e engenharia de atributos. A combinação de dados "
        "geográficos com duas datas de satélite permite avaliar presença e variação do sinal "
        "de vegetação sem criar observações artificiais.",
    )

    doc.add_heading("2 Pergunta e objetivos", level=1)
    adicionar_paragrafo(
        doc,
        "Pergunta de Ciência de Dados: quais pontos da BR-116 apresentam maior presença ou "
        "aumento de vegetação no entorno imediato e devem ser priorizados para inspeção?",
        negrito_inicial="Pergunta de Ciência de Dados:",
    )
    adicionar_paragrafo(
        doc,
        "O objetivo foi construir um pipeline capaz de obter cenas reais, preparar dados, "
        "calcular indicadores de vegetação, criar rótulos explicáveis e apresentar os "
        "resultados em tabela e gráficos.",
    )
    adicionar_lista(
        doc,
        [
            "Obter e documentar fontes públicas da rodovia e das imagens de satélite.",
            "Gerar pontos espacialmente distribuídos nas pistas principais.",
            "Calcular NDVI e atributos de qualidade em períodos comparáveis.",
            "Criar uma prioridade reproduzível e discutir seus limites de uso.",
        ],
    )

    doc.add_heading("3 Fontes de dados", level=1)
    adicionar_paragrafo(
        doc,
        "O eixo auxiliar da BR-116 foi obtido do OpenStreetMap por consulta automatizada ao "
        "Overpass. O Sistema Nacional de Viação do DNIT serviu como referência institucional. "
        "O trecho faz parte da Via Dutra administrada pela CCR RioSP, empresa Motiva.",
    )
    adicionar_paragrafo(
        doc,
        "As imagens Sentinel-2 L2A foram consultadas na API STAC do Microsoft Planetary "
        "Computer. O Sentinel-2 possui bandas visíveis e de infravermelho próximo com "
        "resolução de 10 m, adequadas ao cálculo do NDVI. A API e os ativos usados não "
        "exigiram serviço pago.",
    )
    tabela(
        doc,
        ["Data", "Tile", "Cobertura de nuvens"],
        [
            ["23 de agosto de 2024", "23KLQ", "0,002 por cento"],
            ["23 de agosto de 2024", "23KMQ", "0,003 por cento"],
            ["13 de agosto de 2025", "23KLQ", "0,015 por cento"],
            ["13 de agosto de 2025", "23KMQ", "1,887 por cento"],
        ],
        larguras=[3.0, 1.3, 2.3],
        alinhamentos=[WD_ALIGN_PARAGRAPH.CENTER] * 3,
    )

    doc.add_heading("4 Método", level=1)
    adicionar_paragrafo(
        doc,
        "O traçado foi filtrado pelas pistas classificadas como motorway. O algoritmo "
        "selecionou 92 pontos com distância mínima aproximada de 150 m. Para cada data, "
        "foram lidas as bandas B04, B08 e a classificação da cena SCL apenas na área de "
        "estudo. Pixels sem dados, saturados, com sombra, nuvem, cirrus ou neve foram "
        "excluídos.",
    )
    adicionar_paragrafo(
        doc,
        "O índice de vegetação pela diferença normalizada foi calculado com o vermelho B04 "
        "e o infravermelho próximo B08.",
    )
    adicionar_equacao_ndvi(doc)
    adicionar_paragrafo(
        doc,
        "Em um raio de 60 m ao redor de cada ponto foram calculados média, mediana, "
        "percentil 90, fração de pixels com NDVI igual ou superior a 0,5 e validade dos "
        "pixels. O percentil 90 representa a parte mais vegetada da janela e reduz a "
        "diluição provocada pelo pavimento no centro.",
    )

    doc.add_heading("5 Regra de rotulagem", level=1)
    adicionar_paragrafo(
        doc,
        "Como não há inspeções de campo com classes verdadeiras, foi criado um rótulo proxy "
        "de prioridade. A regra foi aplicada igualmente aos 92 pontos e não usa quantis para "
        "produzir classes balanceadas.",
    )
    adicionar_lista(
        doc,
        [
            "Alta quando o NDVI p90 de 2025 é pelo menos 0,40 ou o aumento do NDVI médio é pelo menos 0,05.",
            "Média quando o NDVI p90 de 2025 é pelo menos 0,30 ou o aumento do NDVI médio é pelo menos 0,025.",
            "Baixa para os demais pontos.",
        ],
    )
    adicionar_paragrafo(
        doc,
        "O rótulo é reproduzível e sua justificativa permanece em cada linha do dataset. "
        "Ele indica onde iniciar uma verificação, mas não informa espécie, altura, "
        "inclinação, invasão do acostamento ou risco de queda.",
    )

    doc.add_heading("6 Dataset construído", level=1)
    adicionar_paragrafo(
        doc,
        "O dataset longo contém 184 linhas, correspondentes a 92 pontos em duas datas. "
        "Todos os pontos tiveram observações comparáveis. A fração média de pixels válidos "
        "foi 100 por cento nos recortes analisados. A base final contém uma linha por ponto, "
        "atributos dos dois anos, variação temporal, índice, classe e critério do rótulo.",
    )
    tabela(
        doc,
        ["Arquivo", "Linhas", "Conteúdo"],
        [
            ["observacoes_sentinel.csv", "184", "Métricas por ponto e período"],
            ["segmentos_priorizados.csv", "92", "Comparação e prioridade por ponto"],
            ["itens_sentinel.json", "4 cenas", "Procedência e cobertura de nuvens"],
        ],
        larguras=[2.5, 1.0, 3.1],
        alinhamentos=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )

    doc.add_heading("7 Resultados", level=1)
    adicionar_paragrafo(
        doc,
        "O NDVI médio geral foi 0,1504 em 2024 e 0,1440 em 2025. A variação média foi "
        "-0,0064 e a mediana -0,0046. Portanto, as duas observações não indicam aumento "
        "generalizado da vegetação no trecho. Essa pequena diferença também pode refletir "
        "umidade, iluminação, atmosfera e fenologia.",
    )
    tabela(
        doc,
        ["Prioridade", "Pontos", "Percentual"],
        [["Baixa", "65", "70,7"], ["Média", "22", "23,9"], ["Alta", "5", "5,4"]],
        larguras=[2.6, 1.8, 2.2],
        alinhamentos=[WD_ALIGN_PARAGRAPH.CENTER] * 3,
    )
    adicionar_figura(
        doc,
        ROOT / "outputs/figures/distribuicao_prioridade.png",
        "Figura 1 Distribuição dos pontos por prioridade de inspeção",
    )
    adicionar_figura(
        doc,
        ROOT / "outputs/figures/comparacao_ndvi.png",
        "Figura 2 Comparação do NDVI médio por ponto em 2024 e 2025",
    )
    adicionar_figura(
        doc,
        ROOT / "outputs/figures/variacao_ndvi.png",
        "Figura 3 Distribuição da variação temporal do NDVI médio",
    )

    doc.add_heading("8 Pontos prioritários", level=1)
    altos = segmentos[segmentos["prioridade_inspecao"] == "alta"].copy()
    linhas_altos = []
    for linha in altos.itertuples():
        motivo = "Aumento" if linha.delta_ndvi_media >= 0.05 else "Presença"
        linhas_altos.append(
            [
                linha.ponto_id,
                f"{linha.longitude:.6f}",
                f"{linha.latitude:.6f}",
                f"{linha.ndvi_p90_atual:.3f}",
                f"{linha.delta_ndvi_media:.3f}",
                motivo,
            ]
        )
    tabela(
        doc,
        ["Ponto", "Longitude", "Latitude", "NDVI p90", "Variação", "Motivo"],
        linhas_altos,
        larguras=[0.7, 1.25, 1.25, 1.0, 1.0, 1.25],
        alinhamentos=[WD_ALIGN_PARAGRAPH.CENTER] * 6,
    )
    adicionar_paragrafo(
        doc,
        "Os pontos P070, P088, P087 e P035 foram classificados principalmente pela "
        "presença atual. O P057 foi classificado pelo aumento de 0,0508 no NDVI médio. "
        "Esses locais são candidatos para a primeira inspeção e não ocorrências confirmadas.",
    )

    doc.add_heading("9 Avaliação e limitações", level=1)
    adicionar_paragrafo(
        doc,
        "Não foi treinado um modelo supervisionado porque não existem rótulos independentes "
        "de campo. Treinar e medir um classificador contra classes derivadas dos mesmos "
        "atributos produziria uma avaliação circular. A prova de conceito foi avaliada pela "
        "completude, validade dos pixels, distribuição das classes, rastreabilidade e "
        "reprodução integral do resultado.",
    )
    adicionar_lista(
        doc,
        [
            "A resolução de 10 m não mostra galhos, placas ou invasão do acostamento.",
            "A janela de 60 m combina pista, vegetação, edificações e terrenos vizinhos.",
            "Duas datas não caracterizam uma tendência de crescimento.",
            "As datas diferem em dez dias e ainda podem refletir condições ambientais distintas.",
            "Pistas paralelas podem gerar pontos com áreas de influência semelhantes.",
            "O rótulo proxy não foi validado por especialista ou inspeção de campo.",
            "A classe alta é pequena e esse desbalanceamento real não foi corrigido.",
        ],
    )

    doc.add_heading("10 Evolução da solução", level=1)
    adicionar_paragrafo(
        doc,
        "Uma aplicação operacional deve integrar fotografias recorrentes, inspeções de "
        "campo, histórico de poda e queda, dados meteorológicos e uma série temporal maior. "
        "O retorno das equipes de conservação deve gerar rótulos independentes. Somente "
        "então será adequado treinar e avaliar um modelo supervisionado.",
    )
    adicionar_paragrafo(
        doc,
        "A coleta pode ser mensal na estação de maior crescimento, trimestral nos demais "
        "períodos e extraordinária após tempestades ou ventos severos. Os resultados devem "
        "alimentar um painel geográfico e permanecer sujeitos à decisão humana.",
    )

    doc.add_heading("11 Conclusão", level=1)
    adicionar_paragrafo(
        doc,
        "A prova de conceito demonstrou que fontes públicas podem sustentar um pipeline "
        "automatizado para organizar inspeções de vegetação. Foram construídas 184 "
        "observações, 92 pontos comparáveis e cinco candidatos de alta prioridade. O "
        "trabalho atende ao objetivo de demonstrar coleta, preparação, rotulagem, análise e "
        "evolução em escala, mas confirma que dados de satélite isolados não diagnosticam "
        "risco. O próximo passo é validar os candidatos em campo.",
    )

    doc.add_heading("Referências", level=1)
    referencias = [
        "EUROPEAN SPACE AGENCY. Sentinel-2 Facts and Figures. Disponível em https://www.esa.int/Applications/Observing_the_Earth/Copernicus/Sentinel-2/Facts_and_figures. Acesso em setembro de 2026.",
        "MICROSOFT. Planetary Computer STAC API. Disponível em https://planetarycomputer.microsoft.com/docs/reference/stac/. Acesso em setembro de 2026.",
        "OPENSTREETMAP CONTRIBUTORS. OpenStreetMap data and copyright. Disponível em https://www.openstreetmap.org/copyright. Acesso em setembro de 2026.",
        "DNIT. Plano Nacional de Viação e Sistema Nacional de Viação. Disponível em https://www.gov.br/dnit/pt-br/assuntos/atlas-e-mapas/pnv-e-snv. Acesso em setembro de 2026.",
        "MOTIVA. CCR RioSP Via Dutra e Rio-Santos. Disponível em https://rodovias.motiva.com.br/riosp/sobre/sobre-riosp/. Acesso em setembro de 2026.",
    ]
    for referencia in referencias:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(6)
        fonte(p.add_run(referencia), tamanho=10)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
